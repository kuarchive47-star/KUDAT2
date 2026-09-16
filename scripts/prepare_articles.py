#!/usr/bin/env python3
"""고대신문 XLSX를 Supabase 적재용 JSONL로 정제한다.

엑셀 파일만 기사 데이터 원천으로 사용한다. 네트워크 요청이나 원문 기사 URL
생성은 하지 않으며 이미지/임베드 경로는 UI가 가져오지 않는 비활성 메타데이터로만
분리한다. 원문 이메일은 개인정보 최소화를 위해 적재하지 않는다.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import uuid
from collections import Counter
from datetime import date, datetime, time
from html import unescape
from pathlib import Path

from lxml import etree, html
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


MEDIA_NAMESPACE = uuid.UUID("167f4cda-3510-49ee-8dff-acde8b4cc3a2")

ALLOWED_TAGS = {
    "p", "div", "figure", "figcaption", "strong", "em", "b", "i", "u",
    "a", "br", "hr", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote",
    "ul", "ol", "li", "table", "thead", "tbody", "tfoot", "tr", "td", "th",
    "span", "sup", "sub"
}
DROP_WITH_CONTENT = {"script", "style", "noscript", "form", "object", "embed", "svg", "math"}
ALLOWED_ATTRS = {
    "td": {"colspan", "rowspan"},
    "th": {"colspan", "rowspan", "scope"},
}
BLOCK_TAGS = {"p", "figcaption", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "li"}


def compact(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value, epoch) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    raw = compact(value)
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    match = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", raw)
    if not match:
        return None
    try:
        return date(*map(int, match.groups()))
    except ValueError:
        return None


def parse_time(value, epoch) -> time:
    if value in (None, ""):
        return time(0, 0)
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted.time().replace(microsecond=0)
        if isinstance(converted, time):
            return converted.replace(microsecond=0)
    raw = compact(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(raw, fmt).time()
        except ValueError:
            pass
    return time(0, 0)


def parse_timestamp(date_value, time_value, epoch) -> str | None:
    parsed_date = parse_date(date_value, epoch)
    if not parsed_date:
        return None
    return datetime.combine(parsed_date, parse_time(time_value, epoch)).isoformat(timespec="seconds")


def nullable_int(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def sanitize_html(raw_html: str, article_id: int) -> tuple[str, str, list[dict]]:
    wrapper = html.fragment_fromstring(raw_html or "", create_parent="div")
    media: list[dict] = []
    position = 0

    # 삭제하기 전에 XLSX HTML에 적힌 이미지/iframe 경로를 비활성 메타데이터로 색인한다.
    for element in list(wrapper.iter()):
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag not in {"img", "iframe"}:
            continue
        source_path = unescape(compact(element.get("src", "")))
        if not source_path:
            continue
        position += 1
        caption = ""
        figure = next((a for a in element.iterancestors("figure")), None)
        if figure is not None:
            cap = figure.find(".//figcaption")
            caption = compact(cap.text_content()) if cap is not None else ""
        media.append({
            "id": str(uuid.uuid5(MEDIA_NAMESPACE, f"{article_id}|{source_path}")),
            "article_id": article_id,
            "media_kind": "image" if tag == "img" else "embed",
            "source_path": source_path,
            "caption": caption or None,
            "alt_text": compact(element.get("alt", "")) or None,
            "position_number": position,
        })

    for element in reversed(list(wrapper.iterdescendants())):
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag in DROP_WITH_CONTENT or tag in {"iframe", "img"}:
            element.drop_tree()
            continue
        if tag not in ALLOWED_TAGS:
            element.drop_tag()
            continue

        allowed = ALLOWED_ATTRS.get(tag, set())
        for attribute in list(element.attrib):
            if attribute.lower() not in allowed:
                del element.attrib[attribute]

        # 기사 본문 링크는 텍스트만 남기며, 클릭/외부 접근 가능한 속성은 두지 않는다.

    body_html = "".join(
        etree.tostring(child, encoding="unicode", method="html") for child in wrapper
    ).strip()
    blocks = []
    for element in wrapper.iter():
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag in BLOCK_TAGS:
            text = compact(element.text_content())
            if text and (not blocks or blocks[-1] != text):
                blocks.append(text)
    if not blocks:
        fallback = compact(wrapper.text_content())
        blocks = [fallback] if fallback else []
    return body_html, "\n\n".join(blocks), media


def load_sections(path: Path) -> tuple[dict, dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload["primary"], payload["secondary"]


def write_jsonl(path: Path, records: list[dict]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--section-map", type=Path, default=Path(__file__).parents[1] / "data/section_map.json")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--include-private", action="store_true", help="비공개 행도 JSONL에 포함하되 RLS에서는 숨깁니다.")
    args = parser.parse_args()

    if not args.xlsx.exists():
        parser.error(f"XLSX 파일을 찾을 수 없습니다: {args.xlsx}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    primary_names, secondary_names = load_sections(args.section_map)

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    headers = [compact(value) for value in next(iterator)]
    required = {"번호", "제목", "내용", "발행일자"}
    missing = required.difference(headers)
    if missing:
        raise SystemExit(f"필수 열 누락: {', '.join(sorted(missing))}")

    articles: list[dict] = []
    media_rows: list[dict] = []
    skipped = Counter()
    section_unknown = Counter()
    html_bytes = 0
    text_bytes = 0

    for source_row_number, values in enumerate(iterator, start=2):
        row = dict(zip(headers, values))
        article_id = nullable_int(row.get("번호"))
        published_on = parse_date(row.get("발행일자"), workbook.epoch)
        if not article_id:
            skipped["missing_id"] += 1
            continue
        if not published_on:
            skipped["invalid_publication_date"] += 1
            continue

        is_public = (
            compact(row.get("상태")).upper() == "E"
            and compact(row.get("외부노출여부")).upper() == "Y"
            and compact(row.get("노출")).upper() == "Y"
        )
        if not is_public and not args.include_private:
            skipped["not_public"] += 1
            continue

        body_html, body_text, media = sanitize_html(str(row.get("내용") or ""), article_id)
        primary_code = compact(row.get("1차섹션")) or None
        secondary_code = compact(row.get("2차섹션")) or None
        if primary_code and primary_code not in primary_names:
            section_unknown[primary_code] += 1
        if secondary_code and secondary_code not in secondary_names:
            section_unknown[secondary_code] += 1

        hero_path = unescape(compact(row.get("대표이미지")))
        if hero_path and not any(item["source_path"] == hero_path for item in media):
            media.insert(0, {
                "id": str(uuid.uuid5(MEDIA_NAMESPACE, f"{article_id}|{hero_path}")),
                "article_id": article_id,
                "media_kind": "image",
                "source_path": hero_path,
                "caption": None,
                "alt_text": None,
                "position_number": 0,
            })

        title = compact(row.get("제목"))
        subtitle = compact(row.get("부제목")) or None
        author_name = compact(row.get("기자명")) or None
        keywords = compact(row.get("키워드")) or None
        primary_name = primary_names.get(primary_code) if primary_code else None
        secondary_name = secondary_names.get(secondary_code) if secondary_code else None
        # 본문은 body_text에 이미 있으므로 검색용 메타데이터에 중복 저장하지 않는다.
        # RPC가 검색 시 body_text/entity_search_text와 결합한다.
        search_document = " ".join(filter(None, [
            title, subtitle, author_name, primary_name, secondary_name, keywords
        ]))
        source_hash = hashlib.sha256(
            "\u241f".join([title, subtitle or "", body_html, published_on.isoformat()]).encode("utf-8")
        ).hexdigest()
        page_number = nullable_int(row.get("면"))
        if page_number == 0:
            page_number = None

        articles.append({
            "id": article_id,
            "source_record_key": f"xlsx:Sheet1:{article_id}",
            "title": title,
            "subtitle": subtitle,
            "author_name": author_name,
            "primary_section_code": primary_code,
            "secondary_section_code": secondary_code,
            "primary_section_name": primary_name,
            "secondary_section_name": secondary_name,
            "article_grade": compact(row.get("등급")) or None,
            "article_type": compact(row.get("기사형태")) or None,
            "view_count": nullable_int(row.get("조회수")),
            "serial_number": nullable_int(row.get("시리얼번호")),
            "used_in_print": compact(row.get("지면사용여부")).upper() == "Y",
            "page_number": page_number,
            "source_status": compact(row.get("상태")) or None,
            "external_visibility": compact(row.get("외부노출여부")) or None,
            "online_offline": compact(row.get("온/오프")) or None,
            "published_on": published_on.isoformat(),
            "registered_at": parse_timestamp(row.get("등록일"), row.get("등록시간"), workbook.epoch),
            "updated_at": parse_timestamp(row.get("수정일"), row.get("수정시간"), workbook.epoch),
            "keywords": keywords,
            "body_html": body_html,
            "body_text": body_text,
            "entity_search_text": "",
            "search_document": search_document,
            "source_hash": source_hash,
            "is_public": is_public,
        })
        media_rows.extend(media)
        html_bytes += len(body_html.encode("utf-8"))
        text_bytes += len(body_text.encode("utf-8"))

        if args.limit and len(articles) >= args.limit:
            break

    workbook.close()
    write_jsonl(args.output_dir / "articles.jsonl", articles)
    write_jsonl(args.output_dir / "article_media.jsonl", media_rows)
    report = {
        "source_xlsx": str(args.xlsx),
        "source_mode": "xlsx_only",
        "network_requests": 0,
        "article_count": len(articles),
        "media_count": len(media_rows),
        "skipped": dict(skipped),
        "unknown_section_codes": dict(section_unknown),
        "body_html_utf8_bytes": html_bytes,
        "body_text_utf8_bytes": text_bytes,
        "include_private": args.include_private,
        "limit": args.limit,
    }
    (args.output_dir / "prepare_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
